"""Document parsing utilities with a structured intermediate representation."""

from __future__ import annotations

import os
import re
import tempfile
from dataclasses import dataclass, field
from importlib import import_module
from io import BytesIO
from typing import Any

from docx import Document as DocxDocument

SUPPORTED_EXTENSIONS = frozenset({".txt", ".md", ".pdf", ".docx", ".doc"})
MAX_FILE_SIZE = 10 * 1024 * 1024  # 10MB

_PAGE_NOISE_RE = re.compile(r"^\s*(?:第\s*)?\d+\s*(?:页|/\s*\d+)?\s*$")
_TOC_RE = re.compile(r"(?:^|\s)(?:目录|contents)\s*$", flags=re.IGNORECASE)
_DISCLAIMER_RE = re.compile(
    r"(?:仅供内部|机密|版权归|all rights reserved|未经授权)",
    flags=re.IGNORECASE,
)
_HEADING_RE = re.compile(r"^(#{1,6})\s+(.+)$")
_LIST_RE = re.compile(r"^(?:[-*+]\s+|\d+[.)]\s+)(.+)$")
_TABLE_RE = re.compile(r"^\|.+\|$")
_LIST_LINE_RE = re.compile(r"^(?P<indent>\s*)(?P<marker>(?:[-*+])|(?:\d+[.)]))\s+(?P<text>.+)$")
_CLAUSE_PATTERNS = [
    re.compile(r"^(?P<numbering>\d+(?:\.\d+)*)[.)]?\s+(?P<title>.+)$"),
    re.compile(r"^(?P<numbering>第\d+条)\s*(?P<title>.+)$"),
    re.compile(r"^(?P<numbering>[一二三四五六七八九十]+、)\s*(?P<title>.+)$"),
    re.compile(r"^(?P<numbering>（[一二三四五六七八九十0-9]+）)\s*(?P<title>.+)$"),
]


@dataclass(frozen=True)
class ParsedBlock:
    type: str
    text: str
    level: int | None = None
    metadata: dict[str, Any] = field(default_factory=dict)


@dataclass(frozen=True)
class ParsedNode:
    node_type: str
    text: str
    level: int | None = None
    numbering: str | None = None
    children: list["ParsedNode"] = field(default_factory=list)
    metadata: dict[str, Any] = field(default_factory=dict)


@dataclass(frozen=True)
class ParsedDocument:
    title: str
    metadata: dict[str, Any] = field(default_factory=dict)
    blocks: list[ParsedBlock] = field(default_factory=list)
    roots: list[ParsedNode] = field(default_factory=list)


def normalize_requirements_plaintext(text: str) -> str:
    """Normalize newlines and trim surrounding whitespace."""
    if not text:
        return ""
    return text.replace("\r\n", "\n").replace("\r", "\n").strip()


def parse_uploaded_document(filename: str, content: bytes) -> tuple[str, str | None]:
    parsed, error = parse_uploaded_document_structured(filename, content)
    if error or parsed is None:
        return "", error
    return render_parsed_document(parsed), None


def parse_uploaded_document_structured(
    filename: str,
    content: bytes,
) -> tuple[ParsedDocument | None, str | None]:
    if len(content) > MAX_FILE_SIZE:
        return None, "文件大小超过限制（最大 10MB）"

    ext = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext not in SUPPORTED_EXTENSIONS:
        return None, "不支持的文件格式: %s，支持 %s" % (
            ext,
            ", ".join(sorted(SUPPORTED_EXTENSIONS)),
        )

    try:
        if ext in {".txt", ".md"}:
            parsed = parse_raw_document_content(
                _parse_text(content),
                title=_title_from_filename(filename),
                document_type=ext.lstrip("."),
                filename=filename,
            )
        elif ext == ".pdf":
            parsed = _parse_pdf_document(filename, content)
        elif ext == ".docx":
            parsed = _parse_docx_document(filename, content)
        elif ext == ".doc":
            return None, "当前版本暂不直接支持 .doc，请先转换为 .docx 再上传"
        else:
            return None, "未实现的解析器: %s" % ext
        return clean_parsed_document(parsed), None
    except Exception as exc:  # pragma: no cover - defensive
        return None, "文件解析失败: %s" % str(exc)


def parse_raw_document_content(
    text: str,
    *,
    title: str,
    document_type: str | None = None,
    filename: str | None = None,
) -> ParsedDocument:
    normalized = normalize_requirements_plaintext(text)
    doc_type = (document_type or "txt").lower()
    if doc_type == "md":
        parsed = _parse_markdown_document(normalized, title=title)
    else:
        parsed = _parse_plaintext_document(normalized, title=title)
    metadata = dict(parsed.metadata)
    metadata.update(
        {
            "source_type": "text",
            "filename": filename,
            "document_type": doc_type,
            "parser_name": metadata.get("parser_name") or f"local-{doc_type}",
        }
    )
    return clean_parsed_document(
        ParsedDocument(
            title=parsed.title,
            metadata=metadata,
            blocks=list(parsed.blocks),
            roots=list(parsed.roots),
        )
    )


def clean_parsed_document(document: ParsedDocument) -> ParsedDocument:
    noise_counts: dict[str, int] = {}
    normalized_blocks: list[ParsedBlock] = []

    for block in document.blocks:
        text = _normalize_block_text(block)
        if not text:
            continue
        compact = re.sub(r"\s+", " ", text).strip()
        if len(compact) <= 60:
            noise_counts[compact] = noise_counts.get(compact, 0) + 1
        normalized_blocks.append(
            ParsedBlock(
                type=block.type,
                text=text,
                level=block.level,
                metadata=dict(block.metadata or {}),
            )
        )

    repeated_noise = {
        item
        for item, count in noise_counts.items()
        if count >= 3 and (len(item) <= 30 or _PAGE_NOISE_RE.match(item))
    }

    cleaned: list[ParsedBlock] = []
    previous_key: tuple[str, str] | None = None
    for block in normalized_blocks:
        compact = re.sub(r"\s+", " ", block.text).strip()
        if compact in repeated_noise:
            continue
        if _PAGE_NOISE_RE.match(compact):
            continue
        if block.type == "paragraph" and _TOC_RE.search(compact):
            continue
        if block.type == "paragraph" and _DISCLAIMER_RE.search(compact) and len(compact) <= 80:
            continue
        block_key = (block.type, compact)
        if previous_key == block_key:
            continue
        previous_key = block_key
        cleaned.append(block)

    metadata = dict(document.metadata)
    metadata.setdefault("parse_warnings", [])
    roots = _build_tree_from_blocks(cleaned)
    return ParsedDocument(title=document.title, metadata=metadata, blocks=cleaned, roots=roots)


def render_parsed_document(document: ParsedDocument) -> str:
    if document.roots:
        parts = [_render_node(node) for node in document.roots]
        return "\n\n".join(part for part in parts if part.strip()).strip()
    parts = [_render_block(block) for block in document.blocks]
    return "\n\n".join(part for part in parts if part.strip()).strip()


def _title_from_filename(filename: str) -> str:
    if "." in filename:
        name = filename.rsplit(".", 1)[0].strip()
        if name:
            return name
    return filename or "Untitled document"


def _normalize_block_text(block: ParsedBlock) -> str:
    text = normalize_requirements_plaintext(block.text)
    if not text:
        return ""
    if block.type == "paragraph":
        text = _merge_soft_wrapped_lines(text)
    if block.type == "list":
        return "\n".join(line.rstrip() for line in text.split("\n") if line.strip()).rstrip()
    if block.type == "table":
        lines = [re.sub(r"\s+", " ", line).strip() for line in text.split("\n")]
        return "\n".join(line for line in lines if line)
    return "\n".join(line.rstrip() for line in text.split("\n")).strip()


def _merge_soft_wrapped_lines(text: str) -> str:
    pieces = [line.strip() for line in normalize_requirements_plaintext(text).split("\n")]
    merged: list[str] = []
    buffer = ""
    for piece in pieces:
        if not piece:
            if buffer:
                merged.append(buffer.strip())
                buffer = ""
            continue
        if not buffer:
            buffer = piece
            continue
        if _should_merge_lines(buffer, piece):
            buffer += _soft_join(buffer, piece)
        else:
            merged.append(buffer.strip())
            buffer = piece
    if buffer:
        merged.append(buffer.strip())
    return "\n\n".join(item for item in merged if item)


def _should_merge_lines(previous: str, current: str) -> bool:
    if not previous or not current:
        return False
    if previous.endswith((".", "!", "?", "。", "！", "？", "：", ":", ";", "；")):
        return False
    if _looks_like_clause_line(previous) or _looks_like_clause_line(current):
        return False
    if _LIST_RE.match(current) or _TABLE_RE.match(current):
        return False
    if len(previous) >= 120:
        return False
    return True


def _soft_join(previous: str, current: str) -> str:
    if previous[-1].isascii() and current[0].isascii():
        return f" {current}"
    return current


def _render_block(block: ParsedBlock) -> str:
    if block.type == "heading":
        level = max(1, min(block.level or 1, 6))
        return f"{'#' * level} {block.text.strip()}"
    if block.type == "list":
        items = [line.rstrip() for line in block.text.split("\n") if line.strip()]
        normalized = []
        for item in items:
            stripped = item.lstrip()
            match = _LIST_RE.match(stripped)
            value = match.group(1).strip() if match else stripped
            indent = " " * (len(item) - len(stripped))
            normalized.append(f"{indent}- {value}")
        return "\n".join(normalized)
    if block.type == "code":
        return f"```\n{block.text.strip()}\n```"
    return block.text.strip()


def _render_node(node: ParsedNode, *, indent: int = 0) -> str:
    if node.node_type == "heading":
        level = max(1, min(node.level or 1, 6))
        title = f"{'#' * level} {node.text.strip()}".strip()
        children = [_render_node(child, indent=indent) for child in node.children]
        parts = [title, *[item for item in children if item.strip()]]
        return "\n\n".join(parts).strip()

    if node.node_type == "clause":
        prefix = f"{node.numbering} " if node.numbering else ""
        title = f"{prefix}{node.text}".strip()
        children = [_render_node(child, indent=indent) for child in node.children]
        parts = [title, *[item for item in children if item.strip()]]
        return "\n\n".join(parts).strip()

    if node.node_type == "list":
        return "\n".join(_render_node(child, indent=indent) for child in node.children if child.text or child.children)

    if node.node_type == "list_item":
        line = f"{'  ' * indent}- {node.text}".rstrip()
        child_lines = [_render_node(child, indent=indent + 1) for child in node.children]
        parts = [line, *[item for item in child_lines if item.strip()]]
        return "\n".join(parts).strip()

    if node.node_type == "table":
        header = str(node.metadata.get("header") or "").strip()
        separator = str(node.metadata.get("separator") or "").strip()
        rows = [child.text.strip() for child in node.children if child.text.strip()]
        lines = [line for line in [header, separator, *rows] if line]
        return "\n".join(lines).strip()

    if node.node_type == "code":
        return f"```\n{node.text.strip()}\n```"

    return node.text.strip()


def _parse_text(content: bytes) -> str:
    for encoding in ("utf-8", "gbk", "gb2312", "latin-1"):
        try:
            return content.decode(encoding).strip()
        except UnicodeDecodeError:
            continue
    raise ValueError("无法识别文件编码")


def _parse_pdf_document(filename: str, content: bytes) -> ParsedDocument:
    return _parse_with_docling(filename, content)


def _parse_docx_document(filename: str, content: bytes) -> ParsedDocument:
    try:
        document = DocxDocument(BytesIO(content))
        blocks = _docx_blocks_from_document(document)
        return ParsedDocument(
            title=_title_from_filename(filename),
            metadata={
                "source_type": "file",
                "filename": filename,
                "document_type": "docx",
                "parser_name": "python-docx",
            },
            blocks=blocks,
            roots=_build_tree_from_blocks(blocks),
        )
    except Exception:
        return _parse_with_docling(filename, content)


def _parse_with_docling(filename: str, content: bytes) -> ParsedDocument:
    ext = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext not in {".pdf", ".docx"}:
        raise ValueError(f"Docling parser does not support extension: {ext}")

    try:
        converter_module = import_module("docling.document_converter")
    except ImportError as exc:  # pragma: no cover - environment dependent
        raise ImportError("Docling is required for PDF/DOCX parsing. Please install docling.") from exc

    temp_path = ""
    try:
        with tempfile.NamedTemporaryFile(delete=False, suffix=ext) as handle:
            handle.write(content)
            temp_path = handle.name

        converter = converter_module.DocumentConverter()
        result = converter.convert(temp_path)
        document = getattr(result, "document", result)

        markdown = None
        for attr in ("export_to_markdown", "to_markdown"):
            candidate = getattr(document, attr, None)
            if callable(candidate):
                markdown = candidate()
                break

        if not isinstance(markdown, str) or not markdown.strip():
            if ext == ".pdf":
                raise ValueError("Current version does not support OCR; upload a text-based PDF.")
            raise ValueError("Docling returned empty content for the uploaded document.")

        parsed = _parse_markdown_document(markdown, title=_title_from_filename(filename))
        metadata = dict(parsed.metadata)
        metadata.update(
            {
                "source_type": "file",
                "filename": filename,
                "document_type": ext.lstrip("."),
                "parser_name": "docling",
            }
        )
        return ParsedDocument(
            title=parsed.title,
            metadata=metadata,
            blocks=list(parsed.blocks),
            roots=list(parsed.roots),
        )
    finally:
        if temp_path and os.path.exists(temp_path):
            os.unlink(temp_path)


def _parse_plaintext_document(text: str, *, title: str) -> ParsedDocument:
    blocks = [
        ParsedBlock(type="paragraph", text=paragraph)
        for paragraph in _split_paragraphs(text)
        if paragraph.strip()
    ]
    return ParsedDocument(
        title=title,
        metadata={"parser_name": "local-txt"},
        blocks=blocks,
        roots=_build_tree_from_blocks(blocks),
    )


def _parse_markdown_document(text: str, *, title: str) -> ParsedDocument:
    blocks: list[ParsedBlock] = []
    lines = normalize_requirements_plaintext(text).split("\n")
    paragraph_lines: list[str] = []
    list_lines: list[str] = []
    code_lines: list[str] = []
    table_lines: list[str] = []
    in_code = False

    def flush_paragraph() -> None:
        if paragraph_lines:
            blocks.append(ParsedBlock(type="paragraph", text="\n".join(paragraph_lines).strip()))
            paragraph_lines.clear()

    def flush_list() -> None:
        if list_lines:
            blocks.append(ParsedBlock(type="list", text="\n".join(list_lines).rstrip()))
            list_lines.clear()

    def flush_code() -> None:
        if code_lines:
            blocks.append(ParsedBlock(type="code", text="\n".join(code_lines).rstrip()))
            code_lines.clear()

    def flush_table() -> None:
        if table_lines:
            blocks.append(ParsedBlock(type="table", text="\n".join(table_lines).strip()))
            table_lines.clear()

    for raw_line in lines:
        stripped = raw_line.rstrip()
        if stripped.startswith("```"):
            flush_paragraph()
            flush_list()
            flush_table()
            if in_code:
                flush_code()
            in_code = not in_code
            continue

        if in_code:
            code_lines.append(stripped)
            continue

        if not stripped.strip():
            flush_paragraph()
            flush_list()
            flush_table()
            continue

        heading = _HEADING_RE.match(stripped.strip())
        if heading:
            flush_paragraph()
            flush_list()
            flush_table()
            blocks.append(
                ParsedBlock(
                    type="heading",
                    text=heading.group(2).strip(),
                    level=len(heading.group(1)),
                )
            )
            continue

        if _TABLE_RE.match(stripped.strip()):
            flush_paragraph()
            flush_list()
            table_lines.append(stripped.strip())
            continue
        flush_table()

        if _LIST_LINE_RE.match(stripped) and not _looks_like_clause_line(stripped):
            flush_paragraph()
            list_lines.append(stripped)
            continue

        flush_list()
        paragraph_lines.append(stripped.strip())

    flush_paragraph()
    flush_list()
    flush_table()
    flush_code()

    return ParsedDocument(
        title=title,
        metadata={"parser_name": "local-md"},
        blocks=blocks,
        roots=_build_tree_from_blocks(blocks),
    )


def _split_paragraphs(text: str) -> list[str]:
    if not text.strip():
        return []
    return [
        paragraph.strip()
        for paragraph in re.split(r"\n\s*\n+", normalize_requirements_plaintext(text))
        if paragraph.strip()
    ]


def _build_tree_from_blocks(blocks: list[ParsedBlock]) -> list[ParsedNode]:
    roots: list[ParsedNode] = []
    heading_stack: list[ParsedNode] = []
    clause_stack: list[tuple[int, ParsedNode]] = []

    for block in blocks:
        if block.type == "heading":
            node = ParsedNode(node_type="heading", text=block.text.strip(), level=block.level)
            while heading_stack and int(heading_stack[-1].level or 1) >= int(block.level or 1):
                heading_stack.pop()
            if heading_stack:
                parent = heading_stack[-1]
                parent.children.append(node)
            else:
                roots.append(node)
            heading_stack.append(node)
            clause_stack = []
            continue

        target_children = heading_stack[-1].children if heading_stack else roots
        nodes = _nodes_from_block(block)
        for node in nodes:
            if node.node_type == "clause":
                clause_level = _clause_depth(node.numbering)
                while clause_stack and clause_stack[-1][0] >= clause_level:
                    clause_stack.pop()
                if clause_stack:
                    clause_stack[-1][1].children.append(node)
                else:
                    target_children.append(node)
                clause_stack.append((clause_level, node))
                continue

            if clause_stack:
                clause_stack[-1][1].children.append(node)
            else:
                target_children.append(node)

    return roots


def _nodes_from_block(block: ParsedBlock) -> list[ParsedNode]:
    if block.type == "paragraph":
        clause_node = _clause_node_from_text(block.text)
        if clause_node is not None:
            return [clause_node]
        return [ParsedNode(node_type="paragraph", text=block.text.strip())]
    if block.type == "list":
        return [_build_list_node(block.text)]
    if block.type == "table":
        return [_build_table_node(block.text)]
    if block.type == "code":
        return [ParsedNode(node_type="code", text=block.text.rstrip())]
    return [ParsedNode(node_type=block.type, text=block.text.strip(), level=block.level)]


def _clause_node_from_text(text: str) -> ParsedNode | None:
    lines = [line.strip() for line in normalize_requirements_plaintext(text).split("\n") if line.strip()]
    if not lines:
        return None
    first = lines[0]
    for pattern in _CLAUSE_PATTERNS:
        match = pattern.match(first)
        if not match:
            continue
        numbering = str(match.group("numbering") or "").strip().rstrip(".")
        title = str(match.group("title") or "").strip()
        children: list[ParsedNode] = []
        body = "\n".join(lines[1:]).strip()
        if body:
            children.append(ParsedNode(node_type="paragraph", text=body))
        return ParsedNode(
            node_type="clause",
            text=title,
            numbering=numbering,
            level=_clause_depth(numbering),
            children=children,
        )
    return None


def _looks_like_clause_line(text: str) -> bool:
    stripped = text.strip()
    if stripped.startswith(("-", "*", "+")):
        return False
    return any(pattern.match(stripped) for pattern in _CLAUSE_PATTERNS)


def _clause_depth(numbering: str | None) -> int:
    value = str(numbering or "").strip()
    if not value:
        return 1
    if value.startswith("第") and value.endswith("条"):
        return 1
    if value.endswith("、") or value.startswith("（"):
        return 1
    return len([part for part in value.split(".") if part]) or 1


def _build_list_node(text: str) -> ParsedNode:
    roots: list[ParsedNode] = []
    stack: list[tuple[int, ParsedNode]] = []
    for line in text.split("\n"):
        if not line.strip():
            continue
        match = _LIST_LINE_RE.match(line.rstrip())
        if not match:
            continue
        indent = _indent_level(match.group("indent"))
        node = ParsedNode(
            node_type="list_item",
            text=str(match.group("text") or "").strip(),
            level=indent,
            metadata={"marker": str(match.group("marker") or "").strip()},
        )
        while stack and stack[-1][0] >= indent:
            stack.pop()
        if stack:
            stack[-1][1].children.append(node)
        else:
            roots.append(node)
        stack.append((indent, node))
    return ParsedNode(node_type="list", text="", children=roots)


def _indent_level(indent: str) -> int:
    expanded = indent.replace("\t", "    ")
    return max(0, len(expanded) // 2)


def _build_table_node(text: str) -> ParsedNode:
    rows = [row.strip() for row in text.split("\n") if row.strip()]
    header = rows[0] if rows else ""
    separator = rows[1] if len(rows) > 1 else ""
    data_rows = rows[2:] if len(rows) > 2 else rows[1:]
    children = [
        ParsedNode(node_type="table_row", text=row, metadata={"header": header})
        for row in data_rows
    ]
    return ParsedNode(
        node_type="table",
        text=header,
        children=children,
        metadata={"header": header, "separator": separator},
    )


def _docx_blocks_from_document(document: DocxDocument) -> list[ParsedBlock]:
    blocks: list[ParsedBlock] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style_name = str(getattr(paragraph.style, "name", "") or "")
        if style_name.startswith("Heading"):
            digits = re.findall(r"\d+", style_name)
            level = int(digits[0]) if digits else 1
            blocks.append(ParsedBlock(type="heading", text=text, level=level))
            continue
        blocks.append(ParsedBlock(type="paragraph", text=text))

    for table in document.tables:
        table_lines: list[str] = []
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells]
            if not any(values):
                continue
            table_lines.append("| " + " | ".join(values) + " |")
        if table_lines:
            if len(table_lines) > 1:
                separator = "| " + " | ".join(["---"] * len(table.rows[0].cells)) + " |"
                table_lines.insert(1, separator)
            blocks.append(ParsedBlock(type="table", text="\n".join(table_lines)))
    return blocks
